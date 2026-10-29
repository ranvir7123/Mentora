"""A lesson's starting workbook, built from the lesson's own data.

The lesson file says what the sheet holds at the start (values, formulas, bold cells, number formats,
column widths). `build_starter` writes it as an .xlsx with openpyxl, so no Excel is needed to make it.
`fits_width` keeps every sheet inside the columns Excel shows at 920 px wide, beside the 360 px panel
on a 1280 px screen: nothing a step needs may sit off to the side.

Starter data:
    {"sheets": [{"name": "Budget", "cells": {"A1": "Item", "B8": "=SUM(B2:B6)"},
                 "bold": ["A1:C1"], "formats": {"B2:B8": "0.00"}, "widths": {"A": 14}}],
     "active": "Budget"}
Only "name" is required for a sheet; "active" names the sheet the learner starts on.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Any, Mapping

import openpyxl
from openpyxl.styles import Font

from app.excel.address import col_index, is_address, parse_area

PANEL_FREE_WIDTH_PX = 920
ROW_HEADER_PX = 30
SCROLL_BAR_PX = 18
DEFAULT_WIDTH_CHARS = 8.43
MAX_SHEET_NAME = 31
BAD_NAME_CHARS = set(':\\/?*[]')


def column_px(chars: float) -> float:
    """A column's width on screen at 100% zoom, for a width in characters (Excel's default font)."""
    return chars * 7 + 5


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _sheet_problems(i: int, sheet: Any) -> list[str]:
    if not isinstance(sheet, Mapping):
        return [f"starter: sheet {i + 1} should be an object"]
    name = sheet.get("name")
    if not isinstance(name, str) or not name.strip():
        return [f"starter: sheet {i + 1} has no name"]
    at = f"starter: sheet {name}"
    problems = []
    if len(name) > MAX_SHEET_NAME:
        problems.append(f"{at}: its name is {len(name)} characters; Excel allows {MAX_SHEET_NAME}")
    if set(name) & BAD_NAME_CHARS or name.startswith("'") or name.endswith("'"):
        problems.append(f"{at}: Excel does not allow : \\ / ? * [ ] or a leading or trailing ' in a sheet name")
    cells = sheet.get("cells", {})
    if not isinstance(cells, Mapping):
        problems.append(f"{at}: cells should be addresses and values")
        cells = {}
    for addr, value in cells.items():
        if not (is_address(addr) and ":" not in addr and "!" not in addr):
            problems.append(f"{at}: {addr} is not a cell address like B7")
        elif not (_is_number(value) or isinstance(value, (str, bool))):
            problems.append(f"{at}: {addr} holds {value!r}, which is not a number, text or true/false")
    bold = sheet.get("bold", [])
    if not isinstance(bold, list):
        problems.append(f"{at}: bold should be a list of ranges")
        bold = []
    for rng in bold:
        if not (is_address(rng) and "!" not in rng):
            problems.append(f"{at}: bold has {rng!r}, which is not a range like A1:C1")
    formats = sheet.get("formats", {})
    if not isinstance(formats, Mapping):
        problems.append(f"{at}: formats should be ranges and format strings")
        formats = {}
    for rng, fmt in formats.items():
        if not (is_address(rng) and "!" not in rng):
            problems.append(f"{at}: formats has {rng!r}, which is not a range like B2:B8")
        elif not isinstance(fmt, str) or not fmt:
            problems.append(f"{at}: the format for {rng} should be text like 0.00")
    widths = sheet.get("widths", {})
    if not isinstance(widths, Mapping):
        problems.append(f"{at}: widths should be column letters and widths")
        widths = {}
    for col, chars in widths.items():
        if not (isinstance(col, str) and col.isalpha() and len(col) <= 3):
            problems.append(f"{at}: widths has {col!r}, which is not a column letter")
        elif not (_is_number(chars) and 0 < chars <= 255):
            problems.append(f"{at}: column {col}'s width should be a number from 0 to 255")
    return problems


def _shape_problems(data: Any) -> list[str]:
    if not isinstance(data, Mapping):
        return ["starter: should be an object with sheets and active"]
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        return ["starter: sheets should be a list with at least one sheet"]
    problems = []
    seen: set[str] = set()
    for i, sheet in enumerate(sheets):
        problems += _sheet_problems(i, sheet)
        name = sheet.get("name") if isinstance(sheet, Mapping) else None
        if isinstance(name, str):
            if name.casefold() in seen:
                problems.append(f"starter: the sheet name {name} is used twice")
            seen.add(name.casefold())
    active = data.get("active")
    if not isinstance(active, str) or active.casefold() not in seen:
        problems.append(f"starter: the active sheet {active!r} is not one of its sheets")
    return problems


def _rightmost_column(sheet: Mapping) -> int:
    used = list(sheet.get("cells", {})) + list(sheet.get("bold", [])) + list(sheet.get("formats", {}))
    return max((parse_area(a).right for a in used), default=0)


def fits_width(data: Mapping, width_px: int = PANEL_FREE_WIDTH_PX) -> bool:
    """Does every sheet's used part show at 100% zoom in a window `width_px` wide (logical px)?"""
    for sheet in data["sheets"]:
        widths = {k.upper(): v for k, v in sheet.get("widths", {}).items()}
        total = ROW_HEADER_PX + SCROLL_BAR_PX
        for col in range(1, _rightmost_column(sheet) + 1):
            letters = next((k for k in widths if col_index(k) == col), None)
            total += column_px(widths[letters] if letters else DEFAULT_WIDTH_CHARS)
        if total > width_px:
            return False
    return True


def validate_starter(data: Mapping) -> list[str]:
    """Every problem with the starter data, in plain words; [] when it is good to build."""
    problems = _shape_problems(data)
    if not problems and not fits_width(data):
        problems.append(f"starter: a sheet is wider than the {PANEL_FREE_WIDTH_PX} px Excel gets beside "
                        "the panel; use fewer or narrower columns")
    return problems


def build_starter(data: Mapping, path: Path) -> Path:
    """Write the starter workbook at `path` (its folder is made if needed) and return `path`."""
    problems = _shape_problems(data)
    if problems:
        raise ValueError("the starter cannot be built: " + "; ".join(problems))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in data["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        for addr, value in sheet.get("cells", {}).items():
            ws[addr.replace("$", "").upper()] = value
        for rng in sheet.get("bold", []):
            for addr in parse_area(rng).cells():
                ws[addr].font = Font(bold=True)
        for rng, fmt in sheet.get("formats", {}).items():
            for addr in parse_area(rng).cells():
                ws[addr].number_format = fmt
        for col, chars in sheet.get("widths", {}).items():
            ws.column_dimensions[col.upper()].width = chars
    names = [s["name"].casefold() for s in data["sheets"]]
    wb.active = names.index(data["active"].casefold())
    wb.calculation.fullCalcOnLoad = True       # openpyxl stores no results; Excel works them out on open
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
