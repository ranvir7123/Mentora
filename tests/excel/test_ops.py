import openpyxl
import pytest

from app.core.types import Call, validate_args
from app.excel import ops
from app.excel.ops import OPS, do
from tests.excel.fakes import BUSY, ComRecorder, com_error

SAMPLES = {
    "set_value": {"cell": "Budget!B7", "value": 100},
    "set_formula": {"cell": "Budget!B8", "formula": "=SUM(B2:B6)"},
    "fill_down": {"source": "Budget!C2", "range": "Budget!C2:C6"},
    "set_format": {"range": "Budget!B2:B8", "kind": "currency"},
    "set_bold": {"range": "Budget!A1:C1"},
    "add_chart": {"chart_type": "pie", "source": "Budget!A1:B6", "title": "Spending"},
    "freeze": {"rows": 1, "cols": 0},
    "filter": {"range": "Budget!A1:C6", "column": "A", "criteria": "=Food"},
    "sort": {"range": "Budget!A1:C6", "column": "B", "order": "desc", "header": True},
    "rename_sheet": {"old": "Budget", "new": "My budget"},
    "select": {"range": "Budget!B2:B6"},
    "clear": {"range": "Budget!B7"},
}


def log_of(kind: str, args: dict, values: dict | None = None) -> list[str]:
    log: list[str] = []
    do(ComRecorder(log, "app"), ComRecorder(log, "wb", values), Call(kind, args))
    return log


SHEET = "wb.Worksheets('Budget')"


def test_every_op_has_a_sample_with_valid_settings_and_runs():
    assert sorted(SAMPLES) == sorted(OPS)
    for name, kind in OPS.items():
        assert kind.name == name and kind.summary
        assert validate_args(kind, SAMPLES[name]) == [], name
        assert log_of(name, SAMPLES[name]), name


def test_an_unknown_op_is_refused():
    with pytest.raises(ValueError, match="paint"):
        log_of("paint", {})


def test_values_formulas_formats_and_bold():
    assert log_of("set_value", SAMPLES["set_value"]) == [f"{SHEET}.Range('B7').Value2 = 100"]
    assert log_of("set_formula", SAMPLES["set_formula"]) == [f"{SHEET}.Range('B8').Formula = '=SUM(B2:B6)'"]
    assert log_of("set_format", SAMPLES["set_format"]) == [f"{SHEET}.Range('B2:B8').NumberFormat = '\"₹\"#,##0'"]
    assert log_of("set_format", {"range": "C2:C6", "kind": "percent"}) == \
        ["wb.ActiveSheet.Range('C2:C6').NumberFormat = '0%'"]
    assert log_of("set_bold", SAMPLES["set_bold"]) == [f"{SHEET}.Range('A1:C1').Font.Bold = True"]
    assert log_of("set_bold", {"range": "A1", "on": False}) == ["wb.ActiveSheet.Range('A1').Font.Bold = False"]
    assert log_of("clear", SAMPLES["clear"]) == [f"{SHEET}.Range('B7').ClearContents()"]


def test_fill_down_right_or_copy():
    assert log_of("fill_down", SAMPLES["fill_down"]) == [f"{SHEET}.Range('C2:C6').FillDown()"]
    assert log_of("fill_down", {"source": "B9", "range": "B9:E9"}) == ["wb.ActiveSheet.Range('B9:E9').FillRight()"]
    assert log_of("fill_down", {"source": "C2", "range": "C3:C6"}) == \
        ["wb.ActiveSheet.Range('C2').Copy(wb.ActiveSheet.Range('C3:C6'))"]


def test_freeze_sets_the_split_on_the_window_and_activates_nothing():
    log = log_of("freeze", {"rows": 1, "cols": 0})
    assert log[-3:] == ["wb.Windows(1).SplitRow = 1", "wb.Windows(1).SplitColumn = 0",
                        "wb.Windows(1).FreezePanes = True"]
    assert log[:5] == ["wb.Windows(1).FreezePanes = False", "wb.Windows(1).SplitRow = 0",
                       "wb.Windows(1).SplitColumn = 0", "wb.Windows(1).ScrollRow = 1",
                       "wb.Windows(1).ScrollColumn = 1"]
    assert not any("Activate" in line for line in log)
    assert log_of("freeze", {"rows": 0, "cols": 0}) ==         ["wb.Windows(1).FreezePanes = False", "wb.Windows(1).SplitRow = 0", "wb.Windows(1).SplitColumn = 0"]


def test_chart_sort_filter_rename_and_select():
    chart = log_of("add_chart", SAMPLES["add_chart"])
    assert chart[0] == f"{SHEET}.Shapes.AddChart2(-1, 5, 322.0, 20.0, 360, 216)"
    assert chart[1] == f"{SHEET}.Shapes.AddChart2(-1, 5, 322.0, 20.0, 360, 216).Chart.SetSourceData({SHEET}.Range('A1:B6'))"
    assert chart[2:] == [f"{SHEET}.Shapes.AddChart2(-1, 5, 322.0, 20.0, 360, 216).Chart.HasTitle = True",
                         f"{SHEET}.Shapes.AddChart2(-1, 5, 322.0, 20.0, 360, 216).Chart.ChartTitle.Text = 'Spending'"]
    assert log_of("sort", SAMPLES["sort"]) == \
        [f"{SHEET}.Range('A1:C6').Sort(Key1={SHEET}.Range('B1'), Order1=2, Header=1, Orientation=1)"]
    assert log_of("sort", {"range": "A2:C6", "column": "a", "order": "asc", "header": False}) == \
        ["wb.ActiveSheet.Range('A2:C6').Sort(Key1=wb.ActiveSheet.Range('A2'), Order1=1, Header=2, Orientation=1)"]
    assert log_of("filter", SAMPLES["filter"]) == [f"{SHEET}.Range('A1:C6').AutoFilter(1, '=Food')"]
    assert log_of("filter", {"range": "B1:C6", "column": "C", "criteria": ">0.1"}) == \
        ["wb.ActiveSheet.Range('B1:C6').AutoFilter(2, '>0.1')"]
    assert log_of("filter", {"range": "A1:C6"}, {"wb.ActiveSheet.AutoFilterMode": False}) == \
        ["wb.ActiveSheet.Range('A1:C6').AutoFilter()"]
    assert log_of("filter", {"range": "A1:C6"}, {"wb.ActiveSheet.AutoFilterMode": True}) == []
    # a column with no criteria shows every row: "*" would hide the numbers, as it matches text only
    assert log_of("filter", {"range": "A1:C6", "column": "B"}) == ["wb.ActiveSheet.Range('A1:C6').AutoFilter(2)"]
    assert log_of("rename_sheet", SAMPLES["rename_sheet"]) == ["wb.Worksheets('Budget').Name = 'My budget'"]
    assert log_of("select", SAMPLES["select"]) == [f"{SHEET}.Activate()", f"{SHEET}.Range('B2:B6').Select()"]


def test_a_quoted_sheet_name_is_unquoted():
    assert log_of("clear", {"range": "'My ''own'' sheet'!A1"}) == \
        ["wb.Worksheets(\"My 'own' sheet\").Range('A1').ClearContents()"]


# ------------------------------------------------------------------ through the session

def test_perform_runs_the_op_on_the_com_thread_then_reads(rig, monkeypatch):
    done = []
    monkeypatch.setattr(ops, "do", lambda app, wb, op: done.append((app, wb, op.kind)))
    r = rig()
    r.com.advance(0)
    fut = r.session.perform(Call("set_bold", {"range": "A1"}))
    assert fut.result(0) is None and done == [(r.app, r.app.books[0], "set_bold")]
    r.com.advance(0)
    assert r.reader.reads == 2


def test_perform_waits_out_a_busy_excel(rig, monkeypatch):
    tries = []

    def busy_twice(app, wb, op):
        tries.append(op.kind)
        if len(tries) < 3:
            raise com_error(BUSY)
    monkeypatch.setattr(ops, "do", busy_twice)
    r = rig()
    fut = r.session.perform(Call("clear", {"range": "A1"}))
    assert not fut.done()
    r.com.advance(1)
    assert fut.result(0) is None and len(tries) == 3


def test_perform_gives_up_on_an_excel_busy_for_too_long(rig, monkeypatch):
    def always_busy(app, wb, op):
        raise com_error(BUSY)
    monkeypatch.setattr(ops, "do", always_busy)
    r = rig()
    fut = r.session.perform(Call("clear", {"range": "A1"}))
    r.com.advance(60)
    with pytest.raises(ops.ExcelBusy, match="busy"):
        fut.result(0)


def test_a_perform_waiting_out_a_busy_excel_fails_when_the_session_closes(rig, monkeypatch):
    def always_busy(app, wb, op):
        raise com_error(BUSY)
    monkeypatch.setattr(ops, "do", always_busy)
    r = rig()
    fut = r.session.perform(Call("clear", {"range": "A1"}))
    r.com.advance(0.1)
    r.session.close()                            # its next try was booked, and now never comes
    with pytest.raises(RuntimeError, match="Excel was closed"):
        fut.result(0)


def test_a_refused_op_fails_its_future(rig, monkeypatch):
    def refuse(app, wb, op):
        raise com_error(-2146827284, "no such sheet")
    monkeypatch.setattr(ops, "do", refuse)
    r = rig()
    with pytest.raises(Exception, match="no such sheet"):
        r.session.perform(Call("clear", {"range": "Nope!A1"})).result(0)


def test_perform_after_close_fails_at_once(rig):
    r = rig()
    r.session.close()
    with pytest.raises(RuntimeError, match="closed"):
        r.session.perform(Call("clear", {"range": "A1"})).result(0)


def test_reset_rebuilds_the_starter_and_reopens_it(rig, tmp_path):
    r = rig()
    r.com.advance(0)
    work = tmp_path / "monthly-budget.xlsx"
    first_book = r.app.books[0]
    r.app.rec.calls.clear()
    work.write_text("the learner's changes", encoding="utf-8")
    fut = r.session.reset()
    assert fut.result(0) is None
    assert r.app.rec.calls == ["sink.close", "Close(SaveChanges=False)", "Open(monthly-budget.xlsx)"]
    assert openpyxl.load_workbook(work)["Budget"]["B2"].value == 12000
    assert r.events.attached_to is r.app.books[0] is not first_book
    r.com.advance(5)
    assert r.closed == 0 and r.reader.reads == 2
    r.events.fire("changed")                     # the new workbook's events still ask for reads
    r.com.advance(0)
    assert r.reader.reads == 3
