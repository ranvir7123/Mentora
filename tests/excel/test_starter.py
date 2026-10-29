from pathlib import Path

import openpyxl

from app.core.lesson import load_lesson
from app.excel.starter import build_starter, fits_width, validate_starter


def _budget_starter():
    return load_lesson(Path("lessons/monthly-budget")).starter


def test_budget_starter_builds_and_reads_back(tmp_path):
    p = build_starter(_budget_starter(), tmp_path / "b.xlsx")
    assert p == tmp_path / "b.xlsx"
    ws = openpyxl.load_workbook(p)["Budget"]
    assert ws["B2"].value == 12000 and ws["A8"].value == "Total"
    assert ws["B8"].value is None


def test_formulas_bold_formats_and_widths_are_written(tmp_path):
    data = {"sheets": [{"name": "S", "cells": {"A1": "Cost", "A2": 5, "A3": "=A2*2"},
                        "bold": ["A1:B1"], "formats": {"A2:A3": "0.00"}, "widths": {"A": 20}},
                       {"name": "Other", "cells": {"B2": "x"}}],
            "active": "Other"}
    wb = openpyxl.load_workbook(build_starter(data, tmp_path / "f.xlsx"))
    ws = wb["S"]
    assert ws["A3"].value == "=A2*2"
    assert ws["A1"].font.bold and ws["B1"].font.bold and not ws["A2"].font.bold
    assert ws["A2"].number_format == "0.00" and ws["A3"].number_format == "0.00"
    assert ws.column_dimensions["A"].width == 20
    assert wb.sheetnames == ["S", "Other"] and wb.active.title == "Other"


def test_wide_sheet_does_not_fit():
    assert not fits_width({"sheets": [{"name": "S", "cells": {"Z1": 1}}], "active": "S"})


def test_budget_starter_fits():
    assert fits_width(_budget_starter())
    # 30 + 18 + (14*7+5) + (12*7+5) + (10*7+5) = 315 px
    assert fits_width(_budget_starter(), width_px=315)
    assert not fits_width(_budget_starter(), width_px=314)


def test_a_bold_range_past_the_cells_counts_as_used():
    data = {"sheets": [{"name": "S", "cells": {"A1": 1}, "bold": ["A1:P1"]}], "active": "S"}
    assert not fits_width(data)


def test_validate_starter_names_problems():
    data = {"sheets": [{"name": "Budget", "cells": {"A1": "x", "B0": 3}},
                       {"name": "Budget", "cells": {}}],
            "active": "Summary"}
    problems = validate_starter(data)
    assert len(problems) == 3, problems
    assert any("Budget" in p and "twice" in p for p in problems)
    assert any("B0" in p for p in problems)
    assert any("Summary" in p for p in problems)


def test_the_budget_starter_has_no_problems():
    assert validate_starter(_budget_starter()) == []


def test_a_broken_shape_is_a_problem_not_a_crash():
    assert validate_starter({"active": "S"}) != []
    assert validate_starter({"sheets": "S", "active": "S"}) != []
    assert validate_starter({"sheets": [{"cells": {}}], "active": "S"}) != []
    assert validate_starter({"sheets": [{"name": "S", "cells": {"A1": [1]}}], "active": "S"}) != []
    assert validate_starter({"sheets": [{"name": "x" * 32}], "active": "x" * 32}) != []
    assert validate_starter({"sheets": [{"name": "a/b"}], "active": "a/b"}) != []
