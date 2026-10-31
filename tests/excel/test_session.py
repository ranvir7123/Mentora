import threading
from pathlib import Path

import openpyxl
import pytest

from app.core.types import Rect
from app.excel.com_thread import ComThread
from app.excel.session import TEARDOWN_BUSY_LIMIT_S, ExcelSession
from app.win import windows
from tests.excel.fakes import FakeApp, FakeEvents, FakeReader


def test_open_builds_the_starter_and_delivers_a_first_snapshot(rig, tmp_path):
    r = rig()
    work = tmp_path / "monthly-budget.xlsx"
    assert r.app.opened == [work]
    assert openpyxl.load_workbook(work)["Budget"]["B2"].value == 12000
    assert r.snapshots == []                     # the first read happens on the COM thread, soon after
    r.com.advance(0)
    assert r.snapshots == [{"read": 1, "book": "monthly-budget.xlsx"}]
    assert r.session.latest() == r.snapshots[-1]
    assert r.session.hwnd() == 0                 # invisible
    assert r.events.attached_to is r.app.books[0]


def test_an_event_is_one_read_and_handlers_never_call_excel(rig):
    r = rig()
    r.com.advance(0)
    for name in ("changed", "changed", "changed"):
        r.events.fire(name)                      # the fake raises if a handler touches Excel
    assert r.reader.reads == 1
    r.com.advance(0)
    assert r.reader.reads == 2 and len(r.snapshots) == 2
    assert r.com.errors == []


def test_closed_workbook_is_reported(rig):
    r = rig()
    r.com.advance(0)
    r.events.fire("before_close")
    r.app.books[0].learner_closes()
    r.com.advance(0.5)
    assert r.closed == 1
    r.com.advance(5)                             # Workbooks.Count is 0 on every later check
    assert r.closed == 1


def test_a_close_the_learner_cancels_is_not_reported(rig):
    r = rig()
    r.com.advance(0)
    r.events.fire("before_close")
    r.app.busy = True                            # the "save your changes?" box is open
    r.com.advance(3)
    r.app.busy = False                           # they pressed Cancel: the workbook stays open
    r.com.advance(5)
    assert r.closed == 0


def test_excel_killed_outright_is_reported_once(rig):
    r = rig()
    r.com.advance(0)
    r.app.dead = True
    r.com.advance(1.1)
    assert r.closed == 1
    r.com.advance(5)
    assert r.closed == 1


def test_excel_with_no_workbooks_left_is_reported_without_an_event(rig):
    r = rig()
    r.com.advance(0)
    r.app.books[0].learner_closes()
    r.com.advance(1.1)
    assert r.closed == 1


def test_reads_stop_once_the_workbook_is_closed(rig):
    r = rig()
    r.com.advance(0)
    r.app.books[0].learner_closes()
    r.com.advance(1.1)
    reads = r.reader.reads
    r.events.fire("changed")
    r.com.advance(5)
    assert r.reader.reads == reads and r.com.pending() == 0


def test_close_tears_down_in_order(rig):
    r = rig()
    r.com.advance(0)
    r.app.rec.calls.clear()
    r.session.close()
    assert r.app.rec.calls == ["sink.close", "DisplayAlerts=False", "Close(SaveChanges=False)", "Quit"]
    assert r.closed == 0                         # closing from the program is not the learner closing
    r.session.close()                            # a second close does nothing
    assert r.app.rec.calls.count("Quit") == 1


def test_close_waits_out_a_short_busy_spell_and_still_closes_in_order(rig):
    r = rig()
    r.com.advance(0)
    r.app.rec.calls.clear()
    r.app.refusals = 3                           # Excel says "not now" three times, then answers
    r.session.close()
    assert r.app.rec.calls == ["sink.close"] + ["DisplayAlerts=False"] * 4 + ["Close(SaveChanges=False)", "Quit"]
    assert r.app.books == [] and r.session.errors == [] and len(r.slept) == 3


def ending(monkeypatch, others=()) -> list:
    """Pretend this Excel is a program that can be ended; returns the handles ended."""
    ended = []
    monkeypatch.setattr(windows, "process_id", lambda hwnd: 99)
    monkeypatch.setattr(windows, "open_process", lambda pid, named=None: 77 if (pid, named) == (99, "excel.exe") else 0)
    monkeypatch.setattr(windows, "end_process", ended.append)
    monkeypatch.setattr(windows, "close_handle", lambda handle: None)
    monkeypatch.setattr(windows, "program_windows", lambda pid: [4242, *others])     # 4242: the session's own
    monkeypatch.setattr(windows, "class_name", lambda hwnd: "XLMAIN")
    return ended


def test_close_ends_this_excel_when_it_stays_busy(rig, monkeypatch):
    ended = ending(monkeypatch)
    r = rig()
    r.com.advance(0)
    r.app.busy = True                            # the learner is still typing in a cell, and has left
    r.session.close()
    assert ended == [77]
    assert sum(r.slept) <= TEARDOWN_BUSY_LIMIT_S + 0.25
    assert {"Close(SaveChanges=False)", "Quit"} <= set(r.app.rec.calls)    # every step was still tried
    assert len(r.session.errors) == 1 and "ended" in r.session.errors[0]


def test_close_leaves_a_busy_excel_running_when_it_shows_other_workbooks(rig, monkeypatch):
    ended = ending(monkeypatch, others=[5151])   # the learner opened a file of their own in it
    r = rig()
    r.com.advance(0)
    r.app.busy = True
    r.session.close()
    assert ended == []
    assert len(r.session.errors) == 1 and "left running" in r.session.errors[0]


def test_close_with_save_keeps_the_work_copy_and_resume_opens_it(rig, tmp_path):
    r = rig()
    r.com.advance(0)
    r.session.close(save=True)
    work = tmp_path / "monthly-budget.xlsx"
    assert "Close(SaveChanges=True)" in r.app.rec.calls
    assert work.read_text(encoding="utf-8") == "saved by the fake"

    again = rig(resume=True)
    assert again.app.opened == [work]
    assert work.read_text(encoding="utf-8") == "saved by the fake"   # not rebuilt
    again.session.close()

    fresh = rig(resume=False)
    assert openpyxl.load_workbook(work)["Budget"]["A8"].value == "Total"   # rebuilt from the starter
    fresh.session.close()


def test_resume_with_no_work_copy_builds_the_starter(rig, tmp_path):
    r = rig(resume=True)
    assert openpyxl.load_workbook(tmp_path / "monthly-budget.xlsx")["Budget"]["B2"].value == 12000
    r.session.close()


def test_an_excel_that_cannot_open_the_file_is_quit_and_the_open_fails(rig):
    app = FakeApp()
    app.fail_open = True
    with pytest.raises(RuntimeError, match="could not be opened"):
        rig(app=app)
    assert "Quit" in app.rec.calls


def test_macros_are_forced_off_before_the_workbook_opens(rig):
    r = rig()
    calls = r.app.rec.calls
    assert calls.index("AutomationSecurity=3") < calls.index("Open(monthly-budget.xlsx)")
    assert r.app.AutomationSecurity == 3


def test_an_excel_that_keeps_macros_on_never_opens_the_workbook(rig):
    app = FakeApp()
    app.keeps_macros_on = True
    with pytest.raises(RuntimeError, match="macros"):
        rig(app=app)
    assert app.opened == []
    assert "Quit" in app.rec.calls


def test_an_invisible_excel_cannot_be_placed(rig):
    r = rig()
    with pytest.raises(ValueError, match="no window"):
        r.session.place(Rect(0, 0, 1840, 1500)).result(1)


def test_with_the_real_com_thread_a_snapshot_arrives(budget, tmp_path):
    app = FakeApp()
    events = FakeEvents(app.rec)
    got = threading.Event()
    s = ExcelSession(budget, tmp_path, visible=False, resume=False,
                     on_snapshot=lambda snap: got.set(), on_closed=lambda: None,
                     excel_factory=lambda: app, com=ComThread("test-session"),
                     attach_events=events.attach, reader=FakeReader(app))
    try:
        assert got.wait(5)
        assert s.latest() == {"read": 1, "book": "monthly-budget.xlsx"}
    finally:
        s.close()
    assert app.rec.calls[-1] == "Quit"


def test_close_from_the_com_thread_itself_does_not_hang(budget, tmp_path):
    app = FakeApp()
    events = FakeEvents(app.rec)
    com = ComThread("test-session-2")
    s = ExcelSession(budget, tmp_path, visible=False, resume=False,
                     on_snapshot=lambda snap: None, on_closed=lambda: None,
                     excel_factory=lambda: app, com=com, attach_events=events.attach, reader=FakeReader(app))
    com.submit(s.close).result(5)                # e.g. a callback on the COM thread closes the session
    com.stop()
    assert app.rec.calls[-1] == "Quit"


def test_excel_is_given_the_full_path_when_the_work_folder_is_relative(rig, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)                  # Excel opens a relative path from its own folder, not ours
    (tmp_path / "prove").mkdir()
    r = rig(workdir=Path("prove"))
    assert r.app.opened == [tmp_path / "prove" / "monthly-budget.xlsx"]
    assert r.app.opened[0].is_absolute()
